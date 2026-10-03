"""Attachment -> text. Never executes attachment content (no macros, no scripts)."""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from pathlib import Path
from typing import Callable

from ..models import ExtractedDocument

MAX_UNCOMPRESSED_BYTES = 300 * 1024 * 1024   # zip-bomb guard for docx/xlsx


def safe_filename(name: str) -> str:
    """Strip any path components / odd characters (path traversal protection)."""
    base = re.split(r"[\\/]", name or "")[-1]
    base = re.sub(r"[^\w.\- ]", "_", base).strip(" .")
    return (base[:120] or "attachment")


def decode_text(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _check_zip(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(i.file_size for i in z.infolist()) > MAX_UNCOMPRESSED_BYTES:
            raise ValueError("archive expands beyond the safety limit")


class DocumentExtractor:
    def __init__(self, max_size_bytes: int = 25 * 1024 * 1024, max_chars: int = 2_000_000) -> None:
        self.max_size_bytes = max_size_bytes
        self.max_chars = max_chars
        self._handlers: dict[str, Callable[[bytes, ExtractedDocument], None]] = {}
        for ext in (".txt", ".log", ".md"):
            self.register(ext, self._plain)
        self.register(".csv", self._csv)
        self.register(".json", self._json)
        self.register(".xml", self._xml)
        self.register(".pdf", self._pdf)
        self.register(".docx", self._docx)
        self.register(".xlsx", self._xlsx)
        self.register(".xls", self._xls)

    def register(self, ext: str, handler: Callable[[bytes, ExtractedDocument], None]) -> None:
        """Add support for another format later."""
        self._handlers[ext.lower()] = handler

    @property
    def supported_extensions(self) -> set[str]:
        return set(self._handlers)

    def extract(self, filename: str, mime_type: str, data: bytes, message_id: str) -> ExtractedDocument:
        name = safe_filename(filename)
        doc = ExtractedDocument(filename=name, mime_type=mime_type, source_message_id=message_id)
        if len(data) > self.max_size_bytes:
            doc.extraction_status = "SKIPPED_TOO_LARGE"
            doc.error_message = f"{len(data)} bytes exceeds limit of {self.max_size_bytes}"
            return doc
        handler = self._handlers.get(Path(name).suffix.lower())
        if handler is None:
            doc.extraction_status = "UNSUPPORTED"
            doc.error_message = f"unsupported file type '{Path(name).suffix}'"
            return doc
        try:
            handler(data, doc)
        except Exception as exc:  # corrupted / hostile files must never crash the pipeline
            doc.extraction_status = "FAILED"
            doc.text = ""
            doc.error_message = f"{type(exc).__name__}: {str(exc)[:200]}"
            return doc
        if len(doc.text) > self.max_chars:
            doc.text = doc.text[: self.max_chars]
            doc.metadata["truncated"] = True
        doc.extraction_status = "SUCCESS"
        return doc

    # handlers --------------------------------------------------------------
    def _plain(self, data: bytes, doc: ExtractedDocument) -> None:
        doc.text = decode_text(data)

    def _csv(self, data: bytes, doc: ExtractedDocument) -> None:
        text = decode_text(data)
        doc.text = text
        doc.metadata["rows"] = sum(1 for _ in csv.reader(io.StringIO(text)))

    def _json(self, data: bytes, doc: ExtractedDocument) -> None:
        doc.text = json.dumps(json.loads(decode_text(data)), indent=2, ensure_ascii=False)

    def _xml(self, data: bytes, doc: ExtractedDocument) -> None:
        from defusedxml import ElementTree  # blocks entity-expansion / external entities
        root = ElementTree.fromstring(data)
        parts: list[str] = []
        for el in root.iter():
            tag = el.tag.split("}")[-1] if isinstance(el.tag, str) else ""
            attrs = " ".join(f'{k}="{v}"' for k, v in el.attrib.items())
            text = (el.text or "").strip()
            if text or attrs:
                parts.append(f"{tag} {attrs} {text}".strip())
        doc.text = "\n".join(parts)

    def _pdf(self, data: bytes, doc: ExtractedDocument) -> None:
        import pymupdf
        with pymupdf.open(stream=data, filetype="pdf") as pdf:
            if pdf.needs_pass:
                raise ValueError("password-protected PDF")
            doc.page_count = pdf.page_count
            parts, total = [], 0
            for page in pdf:
                t = page.get_text()
                parts.append(t)
                total += len(t)
                if total > self.max_chars:
                    break
        doc.text = "\n".join(parts)

    def _docx(self, data: bytes, doc: ExtractedDocument) -> None:
        from docx import Document
        _check_zip(data)
        d = Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        doc.text = "\n".join(parts)

    def _xlsx(self, data: bytes, doc: ExtractedDocument) -> None:
        import openpyxl
        _check_zip(data)
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)  # values only
        try:
            doc.sheet_count = len(wb.sheetnames)
            parts, total = [], 0
            for ws in wb.worksheets:
                parts.append(f"[sheet: {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    line = " | ".join("" if v is None else str(v) for v in row).strip(" |")
                    if line:
                        parts.append(line)
                        total += len(line)
                if total > self.max_chars:
                    break
        finally:
            wb.close()
        doc.text = "\n".join(parts)

    def _xls(self, data: bytes, doc: ExtractedDocument) -> None:
        import pandas as pd
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None, engine="xlrd")
        doc.sheet_count = len(sheets)
        parts = []
        for name, df in sheets.items():
            parts.append(f"[sheet: {name}]")
            parts.append(df.fillna("").astype(str).to_csv(sep="|", index=False, header=False))
        doc.text = "\n".join(parts)
