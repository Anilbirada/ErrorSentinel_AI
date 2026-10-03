import io
import json

import pytest

from app.extraction.documents import DocumentExtractor, safe_filename

X = DocumentExtractor(max_size_bytes=1_000_000)
LINE = "ERR-5021 Database connection timeout"


def ok(name, data):
    d = X.extract(name, "application/octet-stream", data, "m1")
    assert d.extraction_status == "SUCCESS", d.error_message
    assert "ERR-5021" in d.text
    return d


def test_txt():  ok("a.txt", LINE.encode())
def test_log():  ok("a.log", LINE.encode())
def test_md():   ok("a.md", f"# t\n{LINE}".encode())
def test_csv():  assert ok("a.csv", f"id,msg\n1,{LINE}\n".encode()).metadata["rows"] == 2
def test_json(): ok("a.json", json.dumps({"events": [{"msg": LINE}]}).encode())
def test_xml():  ok("a.xml", f"<logs><e level='x'>{LINE}</e></logs>".encode())
def test_utf16_encoding_fallback(): ok("a.txt", LINE.encode("utf-16"))


def test_pdf():
    import pymupdf
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), LINE)
    d = ok("a.pdf", pdf.tobytes())
    assert d.page_count == 1


def test_docx():
    from docx import Document
    doc = Document()
    doc.add_paragraph(LINE)
    buf = io.BytesIO()
    doc.save(buf)
    ok("a.docx", buf.getvalue())


def test_xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["id", LINE])
    buf = io.BytesIO()
    wb.save(buf)
    assert ok("a.xlsx", buf.getvalue()).sheet_count == 1


def test_xls():
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    wb.add_sheet("s").write(0, 0, LINE)
    buf = io.BytesIO()
    wb.save(buf)
    assert ok("a.xls", buf.getvalue()).sheet_count == 1


@pytest.mark.parametrize("name", ["a.pdf", "a.docx", "a.xlsx", "a.xls", "a.json", "a.xml"])
def test_corrupted_files_fail_cleanly(name):
    d = X.extract(name, "x", b"\x00garbage not a real file", "m1")
    assert d.extraction_status == "FAILED" and d.error_message and d.text == ""


def test_unsupported_type():
    assert X.extract("a.exe", "x", b"MZ", "m1").extraction_status == "UNSUPPORTED"


def test_large_attachment_limit():
    small = DocumentExtractor(max_size_bytes=10)
    assert small.extract("a.txt", "x", b"x" * 11, "m1").extraction_status == "SKIPPED_TOO_LARGE"


def test_xml_entity_bomb_rejected():
    bomb = b'<!DOCTYPE x [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;&a;">]><x>&b;</x>'
    assert X.extract("a.xml", "x", bomb, "m1").extraction_status == "FAILED"


def test_safe_filename_blocks_path_traversal():
    assert safe_filename("../../etc/passwd") == "passwd"
    assert safe_filename("..\\..\\win.ini") == "win.ini"
